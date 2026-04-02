#!/usr/bin/env python3
"""Export pointwise prompt/response rows with human and judge scores to Excel."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
CURRENT_SCORE_FIELDS = (
    "writing_quality_score",
    "evidence_use_score",
    "cti_concept_focus_score",
)


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                rows.append(json.loads(line))
    return rows


def load_score_map(path: Path) -> dict[str, dict[str, Any]]:
    score_map: dict[str, dict[str, Any]] = {}
    for row in load_jsonl(path):
        item_id = str(row.get("item_id") or "")
        if not item_id:
            continue
        score_map[item_id] = row
    return score_map


def autosize_columns(ws) -> None:
    for column_cells in ws.columns:
        max_length = 0
        column = column_cells[0].column
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        adjusted = min(max(max_length + 2, 12), 80)
        ws.column_dimensions[get_column_letter(column)].width = adjusted


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--subset", required=True, help="Subset directory name")
    parser.add_argument("--human-file", required=True, help="Human annotation JSONL path")
    parser.add_argument("--judge-file", required=True, help="Judge score JSONL path")
    parser.add_argument("--human-name", default="human", help="Display label prefix for human columns")
    parser.add_argument("--judge-name", default="gpt", help="Display label prefix for judge columns")
    parser.add_argument("--output", help="Output .xlsx path")
    args = parser.parse_args()

    subset_dir = ROOT / args.subset
    blind_path = subset_dir / "blind_responses.jsonl"
    key_path = subset_dir / "key.jsonl"

    if not blind_path.exists():
        raise FileNotFoundError(f"Missing blind responses file: {blind_path}")
    if not key_path.exists():
        raise FileNotFoundError(f"Missing key file: {key_path}")

    output_path = Path(args.output) if args.output else subset_dir / f"{Path(args.human_file).stem}_vs_{Path(args.judge_file).stem}.xlsx"

    blind_rows = load_jsonl(blind_path)
    blind_by_item = {str(row["item_id"]): row for row in blind_rows}
    key_rows = load_jsonl(key_path)
    key_by_item = {str(row["item_id"]): row for row in key_rows}
    human_scores = load_score_map(Path(args.human_file))
    judge_scores = load_score_map(Path(args.judge_file))

    ordered_item_ids = [str(row["item_id"]) for row in blind_rows if str(row["item_id"]) in human_scores and str(row["item_id"]) in judge_scores]

    wb = Workbook()
    ws = wb.active
    ws.title = "scores"

    headers = [
        "item_id",
        "task",
        "task_label",
        "model_name",
        "word_count",
        "prompt",
        "response",
        f"{args.human_name}_writing_quality",
        f"{args.human_name}_evidence_use",
        f"{args.human_name}_cti_concept_focus",
        f"{args.human_name}_total",
        f"{args.judge_name}_writing_quality",
        f"{args.judge_name}_evidence_use",
        f"{args.judge_name}_cti_concept_focus",
        f"{args.judge_name}_total",
    ]
    ws.append(headers)

    for item_id in ordered_item_ids:
        blind = blind_by_item[item_id]
        key = key_by_item.get(item_id, {})
        human = human_scores[item_id]
        judge = judge_scores[item_id]
        ws.append(
            [
                item_id,
                blind.get("task", ""),
                blind.get("task_label", ""),
                key.get("model_name", ""),
                blind.get("word_count", ""),
                blind.get("prompt", ""),
                blind.get("response", ""),
                human.get("writing_quality_score", ""),
                human.get("evidence_use_score", ""),
                human.get("cti_concept_focus_score", ""),
                human.get("total_score", ""),
                judge.get("writing_quality_score", ""),
                judge.get("evidence_use_score", ""),
                judge.get("cti_concept_focus_score", ""),
                judge.get("total_score", ""),
            ]
        )

    header_font = Font(bold=True)
    wrap_alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = wrap_alignment

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap_alignment

    ws.freeze_panes = "A2"
    autosize_columns(ws)
    ws.column_dimensions["F"].width = 80
    ws.column_dimensions["G"].width = 80

    meta = wb.create_sheet("meta")
    meta_rows = [
        ("subset", args.subset),
        ("human_file", str(Path(args.human_file).resolve())),
        ("judge_file", str(Path(args.judge_file).resolve())),
        ("human_name", args.human_name),
        ("judge_name", args.judge_name),
        ("row_count", len(ordered_item_ids)),
        ("score_fields", ", ".join(CURRENT_SCORE_FIELDS)),
    ]
    for row in meta_rows:
        meta.append(row)
    for cell in meta[1]:
        cell.font = header_font
    autosize_columns(meta)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
